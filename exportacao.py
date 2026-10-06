from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NOME_EMPRESA = 'SGDI'
TITULO_RELATORIO = 'Relatório de gerenciamento de demandas'
PRIORIDADE_ROTULOS = {
    0: 'Nenhuma',
    1: 'Baixa',
    2: 'Média',
    3: 'Alta',
}


def descrever_filtros(status, priority, responsible_id, nome_responsavel, status_rotulos):
    """Texto explícito de cada filtro, inclusive quando o valor é Todos."""

    if status in status_rotulos:
        texto_status = status_rotulos[status]
    else:
        texto_status = 'Todos'

    if priority in PRIORIDADE_ROTULOS:
        texto_prioridade = PRIORIDADE_ROTULOS[priority]
    else:
        texto_prioridade = 'Todas'

    if responsible_id:
        texto_responsavel = nome_responsavel or f'Responsável {responsible_id}'
    else:
        texto_responsavel = 'Todos'

    return (
        f'Status: {texto_status} | '
        f'Prioridade: {texto_prioridade} | '
        f'Responsável: {texto_responsavel}'
    )


def montar_relatorio(conn, status, priority, responsible_id, agora):
    from app import STATUS_ROTULOS, calcular_indicadores

    indicadores = calcular_indicadores(
        conn,
        agora=agora,
        status=status,
        priority=priority,
        responsible_id=responsible_id,
    )
    nome_responsavel = None
    if responsible_id:
        usuario = conn.execute(
            'SELECT nome FROM usuarios WHERE id = ?',
            (responsible_id,),
        ).fetchone()
        if usuario is not None:
            nome_responsavel = usuario['nome']

    filtros = descrever_filtros(
        status, priority, responsible_id, nome_responsavel, STATUS_ROTULOS
    )

    return {
        'empresa': NOME_EMPRESA,
        'titulo': TITULO_RELATORIO,
        'data_geracao': agora.strftime('%d/%m/%Y %H:%M'),
        'filtros': filtros,
        'indicadores': indicadores,
    }


def _linhas_indicadores(indicadores):
    tempo = indicadores['tempo_medio'] or 'Sem concluídas com data de conclusão'
    return [
        ('Total de demandas', indicadores['total']),
        ('Abertas', indicadores['abertas']),
        ('Concluídas', indicadores['concluidas']),
        ('Atrasadas', indicadores['atrasadas']),
        ('Críticas', indicadores['criticas']),
        ('Tempo médio de resolução', tempo),
    ]


class _PdfRelatorio(FPDF):

    def __init__(self, relatorio):
        super().__init__()
        self.relatorio = relatorio
        self.compress = False
        self.set_auto_page_break(auto=True, margin=22)

    def header(self):
        self.set_font('Helvetica', 'B', 16)
        self.cell(0, 8, self.relatorio['empresa'], new_x='LMARGIN', new_y='NEXT')
        self.set_font('Helvetica', '', 11)
        self.cell(0, 6, self.relatorio['titulo'], new_x='LMARGIN', new_y='NEXT')
        self.cell(
            0,
            6,
            f"Data de geração: {self.relatorio['data_geracao']}",
            new_x='LMARGIN',
            new_y='NEXT',
        )
        self.ln(2)
        self.set_draw_color(51, 51, 51)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(6)

    def footer(self):
        self.set_y(-16)
        self.set_draw_color(51, 51, 51)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(2)
        self.set_font('Helvetica', '', 8)
        self.set_text_color(50, 50, 50)
        self.cell(
            0,
            5,
            f"Filtros aplicados: {self.relatorio['filtros']}",
            new_x='LMARGIN',
            new_y='NEXT',
        )
        self.set_text_color(0, 0, 0)


def _tabela_pdf(pdf, titulo, cabecalho, linhas, larguras):
    pdf.set_font('Helvetica', 'B', 12)
    pdf.cell(0, 8, titulo, new_x='LMARGIN', new_y='NEXT')
    pdf.set_fill_color(204, 204, 204)
    pdf.set_font('Helvetica', 'B', 10)
    for texto, largura in zip(cabecalho, larguras):
        pdf.cell(largura, 8, texto, border=1, fill=True)
    pdf.ln()
    pdf.set_font('Helvetica', '', 10)
    if not linhas:
        pdf.cell(sum(larguras), 8, 'Nenhum registro', border=1, new_x='LMARGIN', new_y='NEXT')
        pdf.ln(4)
        return
    for linha in linhas:
        for texto, largura in zip(linha, larguras):
            pdf.cell(largura, 8, str(texto), border=1)
        pdf.ln()
    pdf.ln(4)


def gerar_pdf(relatorio):
    pdf = _PdfRelatorio(relatorio)
    pdf.add_page()
    indicadores = relatorio['indicadores']
    largura = 90

    _tabela_pdf(
        pdf,
        'Indicadores',
        ('Indicador', 'Valor'),
        _linhas_indicadores(indicadores),
        (largura, largura),
    )
    _tabela_pdf(
        pdf,
        'Por status',
        ('Status', 'Quantidade'),
        [(item['rotulo'], item['valor']) for item in indicadores['por_status']],
        (largura, largura),
    )
    _tabela_pdf(
        pdf,
        'Por prioridade',
        ('Prioridade', 'Quantidade'),
        [(item['rotulo'], item['valor']) for item in indicadores['por_prioridade']],
        (largura, largura),
    )
    _tabela_pdf(
        pdf,
        'Por responsável',
        ('Responsável', 'Total', 'Abertas', 'Concluídas', 'Atrasadas', 'Críticas'),
        [
            (
                linha['nome'],
                linha['total'],
                linha['abertas'],
                linha['concluidas'],
                linha['atrasadas'],
                linha['criticas'],
            )
            for linha in indicadores['por_responsavel']
        ],
        (50, 26, 26, 30, 28, 26),
    )
    return bytes(pdf.output())


def _estilo_excel():
    borda = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )
    return {
        'empresa': Font(name='Calibri', bold=True, size=16, color='333333'),
        'titulo': Font(name='Calibri', size=12, color='333333'),
        'cabecalho': Font(name='Calibri', bold=True, color='FFFFFF'),
        'preenchimento': PatternFill('solid', fgColor='333333'),
        'secao': Font(name='Calibri', bold=True, size=12),
        'rodape': Font(name='Calibri', italic=True, size=10, color='555555'),
        'borda': borda,
        'centro': Alignment(horizontal='center'),
    }


def _escrever_tabela(planilha, linha, titulo, cabecalho, registros, estilos):
    planilha.cell(linha, 1, titulo).font = estilos['secao']
    linha += 1
    for coluna, texto in enumerate(cabecalho, start=1):
        celula = planilha.cell(linha, coluna, texto)
        celula.font = estilos['cabecalho']
        celula.fill = estilos['preenchimento']
        celula.border = estilos['borda']
        celula.alignment = estilos['centro']
    linha += 1
    if not registros:
        celula = planilha.cell(linha, 1, 'Nenhum registro')
        celula.border = estilos['borda']
        return linha + 2
    for registro in registros:
        for coluna, valor in enumerate(registro, start=1):
            celula = planilha.cell(linha, coluna, valor)
            celula.border = estilos['borda']
            if coluna > 1:
                celula.alignment = estilos['centro']
        linha += 1
    return linha + 1


def gerar_excel(relatorio):
    indicadores = relatorio['indicadores']
    estilos = _estilo_excel()
    livro = Workbook()
    planilha = livro.active
    planilha.title = 'Gerenciamento'

    planilha['A1'] = relatorio['empresa']
    planilha['A1'].font = estilos['empresa']
    planilha['A2'] = relatorio['titulo']
    planilha['A2'].font = estilos['titulo']
    planilha['A3'] = f"Data de geração: {relatorio['data_geracao']}"
    planilha['A3'].font = estilos['titulo']

    linha = _escrever_tabela(
        planilha,
        5,
        'Indicadores',
        ('Indicador', 'Valor'),
        _linhas_indicadores(indicadores),
        estilos,
    )
    linha = _escrever_tabela(
        planilha,
        linha,
        'Por status',
        ('Status', 'Quantidade'),
        [(item['rotulo'], item['valor']) for item in indicadores['por_status']],
        estilos,
    )
    linha = _escrever_tabela(
        planilha,
        linha,
        'Por prioridade',
        ('Prioridade', 'Quantidade'),
        [(item['rotulo'], item['valor']) for item in indicadores['por_prioridade']],
        estilos,
    )
    linha = _escrever_tabela(
        planilha,
        linha,
        'Por responsável',
        ('Responsável', 'Total', 'Abertas', 'Concluídas', 'Atrasadas', 'Críticas'),
        [
            (
                linha_resp['nome'],
                linha_resp['total'],
                linha_resp['abertas'],
                linha_resp['concluidas'],
                linha_resp['atrasadas'],
                linha_resp['criticas'],
            )
            for linha_resp in indicadores['por_responsavel']
        ],
        estilos,
    )

    texto_rodape = f"Filtros aplicados: {relatorio['filtros']}"
    planilha.cell(linha, 1, texto_rodape).font = estilos['rodape']

    planilha.oddHeader.left.text = relatorio['empresa']
    planilha.evenHeader.left.text = relatorio['empresa']
    planilha.oddHeader.right.text = f"Data de geração: {relatorio['data_geracao']}"
    planilha.evenHeader.right.text = f"Data de geração: {relatorio['data_geracao']}"
    planilha.oddFooter.left.text = texto_rodape
    planilha.evenFooter.left.text = texto_rodape
    planilha.page_setup.orientation = 'landscape'
    planilha.page_setup.fitToPage = True
    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 1
    planilha.page_setup.paperSize = planilha.PAPERSIZE_A4
    planilha.sheet_properties.pageSetUpPr.fitToPage = True
    planilha.page_setup.horizontalCentered = True
    planilha.sheet_view.showGridLines = False
    planilha.page_margins.left = 0.6
    planilha.page_margins.right = 0.6
    planilha.page_margins.header = 0.4
    planilha.page_margins.footer = 0.4
    planilha.oddFooter.right.text = 'Página &P de &N'
    planilha.evenFooter.right.text = 'Página &P de &N'

    larguras = (28, 18, 16, 16, 16, 16)
    for indice, largura in enumerate(larguras, start=1):
        planilha.column_dimensions[get_column_letter(indice)].width = largura

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def nome_arquivo(relatorio, extensao):
    carimbo = relatorio['data_geracao'].replace('/', '-').replace(':', '').replace(' ', '-')
    return f'relatorio-gerenciamento-{carimbo}.{extensao}'
