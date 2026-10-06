import os
import sqlite3
import tempfile
import unittest
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app import (
    api_gerenciamento,
    app,
    calcular_indicadores,
    definir_data_conclusao,
    exportar_gerenciamento_excel,
    exportar_gerenciamento_pdf,
    formatar_duracao,
    gerenciamento,
    get_db,
)
from exportacao import descrever_filtros, montar_relatorio
from test_paginacao import salvar_banco


AGORA = datetime(2026, 10, 6, 12, 0, 0)


def inserir(conn, id_demanda, status, priority, responsible_id, data_criacao, data_conclusao=None):
    conn.execute(
        '''
        INSERT INTO demandas (
            id, titulo, descricao, solicitante,
            data_criacao, priority, usuario_id, status,
            responsible_id, data_conclusao
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''',
        (
            id_demanda,
            f'Demanda {id_demanda}',
            'Descricao',
            'Ana Lima',
            data_criacao,
            priority,
            1,
            status,
            responsible_id,
            data_conclusao,
        ),
    )


class TestDefinicaoDosIndicadores(unittest.TestCase):

    def test_data_de_conclusao_so_na_transicao(self):
        self.assertIsNone(definir_data_conclusao('aberta', agora=AGORA))
        self.assertEqual(
            definir_data_conclusao('concluida', agora=AGORA),
            '2026-10-06 12:00:00',
        )

        legado = {'status': 'concluida', 'data_conclusao': None}
        self.assertIsNone(definir_data_conclusao('concluida', legado, AGORA))

        ja_concluida = {
            'status': 'concluida',
            'data_conclusao': '2026-09-01 08:00:00',
        }
        self.assertEqual(
            definir_data_conclusao('aberta', ja_concluida, AGORA),
            None,
        )
        self.assertEqual(
            definir_data_conclusao('concluida', ja_concluida, AGORA),
            '2026-09-01 08:00:00',
        )

    def test_formatacao_do_tempo_medio(self):
        self.assertIsNone(formatar_duracao(None))
        self.assertEqual(formatar_duracao(86400), '1,0 dia')
        self.assertEqual(formatar_duracao(2.5 * 86400), '2,5 dias')
        self.assertEqual(formatar_duracao(3600), '1,0 hora')
        self.assertEqual(formatar_duracao(90), '2 minutos')


class TestCalculoDosIndicadores(unittest.TestCase):

    def setUp(self):
        self.arquivo = tempfile.NamedTemporaryFile(suffix='.db', delete=False)
        self.arquivo.close()
        self.database_original = app.config.get('DATABASE')
        app.config['DATABASE'] = self.arquivo.name
        salvar_banco(self.arquivo.name, 0)
        self.conn = get_db()

    def tearDown(self):
        self.conn.close()
        if self.database_original is None:
            app.config.pop('DATABASE', None)
        else:
            app.config['DATABASE'] = self.database_original
        os.remove(self.arquivo.name)

    def test_banco_vazio(self):
        indicadores = calcular_indicadores(self.conn, AGORA)

        self.assertEqual(indicadores['total'], 0)
        self.assertEqual(indicadores['abertas'], 0)
        self.assertEqual(indicadores['concluidas'], 0)
        self.assertEqual(indicadores['atrasadas'], 0)
        self.assertEqual(indicadores['criticas'], 0)
        self.assertIsNone(indicadores['tempo_medio'])
        self.assertEqual(list(indicadores['por_responsavel']), [])

    def test_contagens_atraso_criticidade_e_media(self):
        inserir(self.conn, 1, 'aberta', 3, 1, '2026-09-01 12:00:00')
        inserir(self.conn, 2, 'em_andamento', 1, 1, '2026-09-06 12:00:00')
        inserir(self.conn, 3, 'em_andamento', 3, 2, '2026-09-06 11:59:59')
        inserir(self.conn, 4, 'aberta', 2, None, '2026-10-01 12:00:00')
        inserir(
            self.conn, 5, 'concluida', 3, 2,
            '2026-10-01 12:00:00', '2026-10-03 12:00:00',
        )
        inserir(
            self.conn, 6, 'concluida', 1, 1,
            '2026-10-01 12:00:00', '2026-10-05 12:00:00',
        )
        inserir(self.conn, 7, 'concluida', 3, 1, '2026-08-01 12:00:00')
        inserir(self.conn, 8, 'cancelada', 3, 2, '2026-01-01 12:00:00')
        self.conn.commit()

        indicadores = calcular_indicadores(self.conn, AGORA)

        self.assertEqual(indicadores['total'], 8)
        self.assertEqual(indicadores['abertas'], 4)
        self.assertEqual(indicadores['status_aberta'], 2)
        self.assertEqual(indicadores['em_andamento'], 2)
        self.assertEqual(indicadores['concluidas'], 3)
        self.assertEqual(indicadores['canceladas'], 1)
        self.assertEqual(indicadores['atrasadas'], 2)
        self.assertEqual(indicadores['criticas'], 2)
        self.assertEqual(indicadores['concluidas_sem_data'], 1)
        self.assertEqual(indicadores['tempo_medio_segundos'], 3 * 86400)
        self.assertEqual(indicadores['tempo_medio'], '3,0 dias')

        por_nome = {
            linha['nome']: linha
            for linha in indicadores['por_responsavel']
        }
        self.assertEqual(por_nome['Ana Lima']['total'], 4)
        self.assertEqual(por_nome['Ana Lima']['abertas'], 2)
        self.assertEqual(por_nome['Ana Lima']['concluidas'], 2)
        self.assertEqual(por_nome['Ana Lima']['atrasadas'], 1)
        self.assertEqual(por_nome['Ana Lima']['criticas'], 1)
        self.assertEqual(por_nome['Bruno Souza']['total'], 3)
        self.assertEqual(por_nome['Bruno Souza']['criticas'], 1)
        self.assertEqual(por_nome['Bruno Souza']['atrasadas'], 1)
        self.assertEqual(por_nome['Sem responsável']['total'], 1)
        self.assertEqual(por_nome['Sem responsável']['abertas'], 1)
        self.assertIsNone(por_nome['Sem responsável']['id'])

    def test_pagina_e_api_exibem_os_numeros(self):
        inserir(self.conn, 1, 'aberta', 3, 1, '2026-09-01 12:00:00')
        inserir(
            self.conn, 2, 'concluida', 1, 1,
            '2026-10-04 12:00:00', '2026-10-05 12:00:00',
        )
        self.conn.commit()

        with app.test_request_context('/gerenciamento'):
            html = gerenciamento()

        self.assertIn('Total de demandas', html)
        self.assertIn('Tempo médio de resolução', html)
        self.assertIn('Por responsável', html)
        self.assertIn('Ana Lima', html)
        self.assertIn('1,0 dia', html)
        self.assertIn('href="/?responsibleId=1"', html)

        with app.test_request_context('/api/gerenciamento'):
            dados = api_gerenciamento().get_json()

        self.assertEqual(dados['total'], 2)
        self.assertEqual(dados['abertas'], 1)
        self.assertEqual(dados['concluidas'], 1)
        self.assertEqual(dados['atrasadas'], 1)
        self.assertEqual(dados['criticas'], 1)
        self.assertEqual(dados['tempoMedioResolucao'], '1,0 dia')
        self.assertEqual(dados['porResponsavel'][0]['nome'], 'Ana Lima')
        self.assertEqual(dados['porStatus'][0]['rotulo'], 'Aberta')
        self.assertEqual(dados['porStatus'][0]['valor'], 1)
        self.assertEqual(dados['porPrioridade'][3]['rotulo'], 'Alta')

    def test_filtro_restringe_numeros_e_graficos(self):
        inserir(self.conn, 1, 'aberta', 3, 1, '2026-09-01 12:00:00')
        inserir(
            self.conn, 2, 'concluida', 1, 2,
            '2026-10-01 12:00:00', '2026-10-03 12:00:00',
        )
        self.conn.commit()

        so_alta = calcular_indicadores(self.conn, AGORA, priority=3)
        self.assertEqual(so_alta['total'], 1)
        self.assertEqual(so_alta['criticas'], 1)
        self.assertEqual(so_alta['por_prioridade'][3]['valor'], 1)
        self.assertEqual(so_alta['por_prioridade'][3]['largura'], 100)
        self.assertEqual(so_alta['por_prioridade'][1]['valor'], 0)
        self.assertEqual(so_alta['por_responsavel'][0]['nome'], 'Ana Lima')

        with app.test_request_context('/gerenciamento?status=concluida'):
            html = gerenciamento()

        self.assertIn('Por status', html)
        self.assertIn('Por prioridade', html)
        self.assertIn('id="grafico-status"', html)
        self.assertIn('value="concluida"', html)
        self.assertIn('selected', html)
        self.assertIn('Exibindo 1 demanda(s) do filtro atual.', html)

        with app.test_request_context(
            '/api/gerenciamento?status=aberta&responsibleId=1&priority=3'
        ):
            dados = api_gerenciamento().get_json()

        self.assertEqual(dados['total'], 1)
        self.assertEqual(dados['abertas'], 1)
        self.assertEqual(dados['concluidas'], 0)
        self.assertEqual(dados['porStatus'][2]['valor'], 0)
        self.assertEqual(dados['porResponsavel'][0]['id'], 1)

    def test_preparar_banco_cria_data_conclusao(self):
        conn = sqlite3.connect(self.arquivo.name)
        colunas = {
            linha[1]
            for linha in conn.execute('PRAGMA table_info(demandas)')
        }
        conn.close()

        self.assertIn('data_conclusao', colunas)


class TestExportacaoDoRelatorio(unittest.TestCase):

    def setUp(self):
        self.arquivo = tempfile.NamedTemporaryFile(suffix='.db', delete=False)
        self.arquivo.close()
        self.database_original = app.config.get('DATABASE')
        app.config['DATABASE'] = self.arquivo.name
        salvar_banco(self.arquivo.name, 0)
        self.conn = get_db()
        inserir(self.conn, 1, 'aberta', 3, 1, '2026-09-01 12:00:00')
        inserir(
            self.conn, 2, 'concluida', 1, 2,
            '2026-10-01 12:00:00', '2026-10-03 12:00:00',
        )
        self.conn.commit()

    def tearDown(self):
        self.conn.close()
        if self.database_original is None:
            app.config.pop('DATABASE', None)
        else:
            app.config['DATABASE'] = self.database_original
        os.remove(self.arquivo.name)

    def test_filtros_explicitos_inclusive_quando_vazios(self):
        from app import STATUS_ROTULOS

        texto = descrever_filtros(None, None, None, None, STATUS_ROTULOS)
        self.assertEqual(
            texto,
            'Status: Todos | Prioridade: Todas | Responsável: Todos',
        )
        texto = descrever_filtros('concluida', 3, 1, 'Ana Lima', STATUS_ROTULOS)
        self.assertEqual(
            texto,
            'Status: Concluída | Prioridade: Alta | Responsável: Ana Lima',
        )

    def test_pdf_e_excel_trazem_empresa_data_e_rodape(self):
        relatorio = montar_relatorio(
            self.conn, 'concluida', None, 2, AGORA
        )
        self.assertEqual(relatorio['empresa'], 'SGDI')
        self.assertEqual(relatorio['data_geracao'], '06/10/2026 12:00')
        self.assertIn('Status: Concluída', relatorio['filtros'])
        self.assertIn('Responsável: Bruno Souza', relatorio['filtros'])
        self.assertIn('Prioridade: Todas', relatorio['filtros'])

        with app.test_request_context(
            '/gerenciamento/exportar.pdf?status=concluida&responsibleId=2'
        ):
            resposta = exportar_gerenciamento_pdf()

        self.assertEqual(resposta.mimetype, 'application/pdf')
        self.assertIn('relatorio-gerenciamento-', resposta.headers['Content-Disposition'])
        resposta.direct_passthrough = False
        pdf = resposta.get_data()
        self.assertTrue(pdf.startswith(b'%PDF'))
        texto_pdf = pdf.decode('latin-1', errors='ignore')
        self.assertIn('SGDI', texto_pdf)
        self.assertIn('Filtros aplicados:', texto_pdf)
        self.assertIn('Data de gera', texto_pdf)
        self.assertIn('Status: Conclu', texto_pdf)
        self.assertIn('Bruno Souza', texto_pdf)
        self.assertIn('Prioridade: Todas', texto_pdf)

        with app.test_request_context('/gerenciamento/exportar.xlsx?priority=3'):
            resposta = exportar_gerenciamento_excel()

        resposta.direct_passthrough = False
        livro = load_workbook(BytesIO(resposta.get_data()))
        planilha = livro.active
        self.assertEqual(planilha['A1'].value, 'SGDI')
        self.assertTrue(planilha['A3'].value.startswith('Data de geração:'))
        rodape = planilha.oddFooter.left.text
        self.assertIn('Filtros aplicados:', rodape)
        self.assertIn('Prioridade: Alta', rodape)
        self.assertIn('Status: Todos', rodape)
        self.assertIn('Responsável: Todos', rodape)
        textos = [
            planilha.cell(linha, 1).value
            for linha in range(1, planilha.max_row + 1)
        ]
        self.assertIn(
            'Filtros aplicados: Status: Todos | Prioridade: Alta | Responsável: Todos',
            textos,
        )
        valores = {
            planilha.cell(linha, 1).value: planilha.cell(linha, 2).value
            for linha in range(1, planilha.max_row + 1)
        }
        self.assertEqual(valores['Total de demandas'], 1)
        self.assertEqual(valores['Abertas'], 1)
        self.assertEqual(valores['Concluídas'], 0)
        nomes = textos
        self.assertIn('Ana Lima', nomes)
        self.assertNotIn('Bruno Souza', nomes)

    def test_pagina_oferece_os_dois_formatos_com_filtro(self):
        with app.test_request_context('/gerenciamento?status=concluida&priority=1'):
            html = gerenciamento()

        self.assertIn('formaction="/gerenciamento/exportar.pdf"', html)
        self.assertIn('formaction="/gerenciamento/exportar.xlsx"', html)
        self.assertIn('form="filtros-gerenciamento"', html)
        self.assertIn('Exportar PDF', html)
        self.assertIn('Exportar Excel', html)
